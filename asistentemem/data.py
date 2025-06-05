from docx import Document
import requests
import json
import pandas as pd

document_text = ""
api_text = ""


def cargar_documento(archivo):
    """Load and process a .docx document"""
    global document_text
    print(f"[DEBUG] Starting document load: {archivo.name}")
    try:
        doc = Document(archivo.name)
        print(f"[DEBUG] Total paragraphs found by python-docx: {len(doc.paragraphs)}")
        
        all_text_parts = []
        seen_text_parts = set() # Use a set for efficient "in" checks

        # Log text of first few raw paragraphs for inspection
        for i, p in enumerate(doc.paragraphs):
            stripped_p_text = p.text.strip()
            if i < 5:  # Log first 5 paragraphs as an example
                print(f"[DEBUG] Raw paragraph {i} text (stripped length): {len(stripped_p_text)}, text (first 50 chars, stripped): '{stripped_p_text[:50]}'")
            if stripped_p_text and stripped_p_text not in seen_text_parts:
                all_text_parts.append(stripped_p_text)
                seen_text_parts.add(stripped_p_text)
        
        print(f"[DEBUG] Total tables found: {len(doc.tables)}")
        for i, table in enumerate(doc.tables):
            print(f"[DEBUG] Processing table {i}")
            for row_idx, row in enumerate(table.rows):
                for cell_idx, cell in enumerate(row.cells):
                    cell_text = cell.text.strip()
                    if cell_text and cell_text not in seen_text_parts:
                        # Log first few unique cell texts for inspection
                        # Simplified logging condition to show a few unique cells added
                        if len(all_text_parts) < 20: # Example: log until 20 unique parts are collected or so
                             print(f"[DEBUG] Adding unique cell text: Table {i}, Row {row_idx}, Cell {cell_idx} (stripped length): {len(cell_text)}, text (first 50 chars): '{cell_text[:50]}'")
                        all_text_parts.append(cell_text)
                        seen_text_parts.add(cell_text)
        
        print(f"[DEBUG] Total unique text parts extracted (paragraphs + table cells): {len(all_text_parts)}")
        document_text = "\n\n".join(all_text_parts) # Join with double newline to separate content from different elements
        print(f"[DEBUG] Final document text length: {len(document_text)}")
        if not document_text:
            print("[WARNING] Document text is empty after processing paragraphs and tables.")
            
    except Exception as e:
        print(f"[ERROR] Failed to load document: {str(e)}")
        document_text = ""
    return [
        {
            "role": "assistant",
            "content": "📄 Documento cargado exitosamente. ¿Qué desea preguntar?",
        }
    ]


def obtener_datos_api(url):
    """Retrieve and process API data"""
    global api_text
    try:
        response = requests.get(url)
        print(f"📡 API solicitada: {url}")
        print(f"📡 Código de respuesta: {response.status_code}")

        if response.status_code == 200:
            api_data = response.json()

            if "result" in api_data and "records" in api_data["result"]:
                records = api_data["result"]["records"]

                df = pd.DataFrame(records)

                # Convertir la columna "Fecha" a datetime
                df["Fecha"] = pd.to_datetime(df["Fecha"])

                # Pivotear la tabla: "CodigoVariable" serán las columnas, "Valor" el contenido
                df_pivot = df.pivot(
                    index="Fecha", columns="CodigoVariable", values="Valor"
                ).reset_index()

                # Cambiar el formato de la fecha a "DD-MM-YYYY"
                df_pivot["Fecha"] = df_pivot["Fecha"].dt.strftime("%d-%m-%Y")

                df_pivot.index.name = "index"

                # Convertir el DataFrame a texto legible
                pivot_text = df_pivot.to_string()

                name = api_data["result"].get("name", "Sin nombre")
                description = (
                    api_data["result"]
                    .get("metadata", {})
                    .get("description", "Sin descripción")
                )

                api_text = (
                    f"🔹 **{name}**\n"
                    f"📄 {description}\n\n"
                    f"🌐 **Información cargada desde la API:** {url}\n\n"
                    f"📊 **Conjuntos de dato del API:**\n```\n{pivot_text}\n```"
                )

                print(f"✅ API cargada: {api_text}")
                return [
                    {
                        "role": "assistant",
                        "content": "📄 API cargada exitosamente. ¿Qué desea preguntar?",
                    }
                ]
            else:
                print("❌ Error: No se encontraron registros en la API.")
                return [
                    {
                        "role": "assistant",
                        "content": "❌ Error: No se encontraron registros en la API.",
                    }
                ]
        else:
            print(f"❌ Error en la API: {response.status_code}")
            return [
                {
                    "role": "assistant",
                    "content": f"❌ Error en la API: {response.status_code}",
                }
            ]
    except Exception as e:
        print(f"❌ Error al conectar con la API: {str(e)}")
        return [
            {
                "role": "assistant",
                "content": f"❌ Error al conectar con la API: {str(e)}",
            }
        ]


def get_document_text():
    """Return the global document_text variable"""
    global document_text
    return document_text


def get_api_text():
    """Return the global api_text variable"""
    global api_text
    return api_text
